"""구매팀 E2E 엑셀(발주관리 시트)을 purchase_lines / purchase_receipts로 가져온다.

앱이 구매 데이터의 원본이고, 엑셀은 전환 기간 동안 다시 올려서 바뀐 줄만 갱신하는 용도다.
- 줄 매칭: 엑셀 관리번호가 파일 안에서 유일하고 DB의 엑셀 출신 줄과 같으면 그 줄. 아니면
  지문(PO번호|제조사PN|품명|발주수량, 같은 지문이 여러 줄이면 #순번)으로 찾는다. 둘 다 없으면 새 줄 —
  관리번호가 비었거나 중복이거나 이미 다른 줄이 쓰고 있으면 새 관리번호(L+일련번호)를 붙인다.
- DB엔 있는데 새 엑셀에 없는 줄은 지우지 않는다 (앱에서 만든 줄일 수도 있어서) — 개수만 알려준다.
- 입고: 앱에서 입고를 등록한 줄은 엑셀의 입고값을 무시한다 (같은 입고를 두 번 세지 않게).
  그 외에는 엑셀의 실제 입고 수량·일자·입고처를 source='excel' 입고 기록 1건으로 맞춘다."""

import io
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pandas as pd
from sqlalchemy import text

from db import get_db_connection

SHEET = "발주관리"
HEADER_ROW = 5

# DB 컬럼 -> 엑셀 헤더(공백/줄바꿈 제거 후). BL번호·수입신고필증 번호는 5행 헤더가 비어 있어서
# '계약 납기' 바로 다음 두 열로 찾는다.
COLUMNS = {
    "excel_line_no": "관리번호",
    "pr_no": "구매요청(PR)번호",
    "pr_requested_on": "요청일자",
    "pr_approved_on": "승인일자",
    "pr_dept": "요청부서",
    "pr_requester": "요청자",
    "pr_basis": "구매요청근거",
    "category": "품목유형",
    "po_no": "발주(PO)번호",
    "po_approved_on": "승인일자#2",
    "buyer": "담당자",
    "supply_type": "내·외자",
    "pi_no": "PI번호(외자)",
    "supplier": "공급사",
    "item_code": "품목번호",
    "item_name": "품명",
    "mfr_pn": "제조사PN(PartNumber)",
    "spec": "규격",
    "qty": "발주수량",
    "unit_price": "단가",
    "currency": "통화",
    "fx_rate": "기준환율",
    "contract_due": "계약납기",
    "received_qty": "실제입고수량",
    "received_on": "실제입고일자",
    "destination": "입고처",
    "excel_status": "납품진행상태",
    "tax_invoice_on": "세금계산서발행일",
    "paid_on": "대금지급일",
    "pay_status": "지급상태",
    "note": "비고",
}
DATE_COLS = {"pr_requested_on", "pr_approved_on", "po_approved_on", "contract_due", "received_on",
             "tax_invoice_on", "paid_on"}
NUM_COLS = {"qty", "unit_price", "fx_rate", "received_qty"}
# purchase_lines에 저장하는(엑셀이 주인인) 필드. 입고 3개 필드는 purchase_receipts로 간다.
LINE_FIELDS = [c for c in COLUMNS if c not in ("excel_line_no", "received_qty", "received_on")] + [
    "bl_no", "import_decl_no", "amount", "amount_krw", "fingerprint",
]
LINE_NO_RE = re.compile(r"^L(\d+)$")


def _key(h):
    return re.sub(r"\s+", "", str(h or ""))


def _clean_text(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s in ("", "0", "-") else s


def _clean_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return pd.Timestamp(str(v).strip().replace(".", "-")).date()
    except (ValueError, TypeError):
        return None


def _clean_num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _locate_columns(header):
    keys = [_key(h) for h in header]
    pos = {}
    seen = Counter()
    for i, k in enumerate(keys):
        if not k:
            continue
        seen[k] += 1
        pos[k if seen[k] == 1 else f"{k}#{seen[k]}"] = i
    missing = [h for h in COLUMNS.values() if h not in pos]
    if missing:
        raise ValueError(f"'{SHEET}' 시트에서 이 열을 찾지 못했어요: {', '.join(missing)}")
    due = pos["계약납기"]
    return {col: pos[h] for col, h in COLUMNS.items()} | {"bl_no": due + 1, "import_decl_no": due + 2}


def read_e2e_excel(file):
    """엑셀(경로/바이트/파일객체)에서 정리된 줄 목록(DataFrame)을 만든다."""
    if isinstance(file, (bytes, bytearray)):
        file = io.BytesIO(file)
    wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError(f"'{SHEET}' 시트가 없어요. 구매 E2E 통합이력 파일이 맞는지 확인해주세요.")
    rows = list(wb[SHEET].iter_rows(min_row=HEADER_ROW, values_only=True))
    wb.close()  # read_only 모드는 닫기 전까지 파일을 잡고 있다
    idx = _locate_columns(rows[0])

    records = []
    for r in rows[1:]:
        get = lambda col: r[idx[col]] if idx[col] < len(r) else None  # noqa: E731
        rec = {}
        for col in idx:
            v = get(col)
            if col in DATE_COLS:
                rec[col] = _clean_date(v)
            elif col in NUM_COLS:
                rec[col] = _clean_num(v)
            else:
                rec[col] = _clean_text(v)
        if rec["qty"] is None and rec["item_name"] is None:
            continue  # 빈 줄
        rec["qty"] = rec["qty"] or 0.0
        if rec["item_code"]:
            rec["item_code"] = rec["item_code"].upper()
        price = rec["unit_price"]
        rec["amount"] = round(rec["qty"] * price, 4) if price is not None else None
        rec["amount_krw"] = (
            round(rec["amount"] * rec["fx_rate"], 2) if rec["amount"] is not None and rec["fx_rate"] else None
        )
        records.append(rec)

    df = pd.DataFrame(records)
    base = (
        df["po_no"].fillna("") + "|" + df["mfr_pn"].fillna("") + "|" + df["item_name"].fillna("")
        + "|" + df["qty"].map(lambda q: f"{q:g}")
    )
    df["fingerprint"] = base + "#" + (base.groupby(base).cumcount() + 1).astype(str)
    return df


def _norm_val(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, (int, float, Decimal)):
        return round(float(v), 6)
    return v


def _params(rec, cols):
    """DB에 넘길 값: NaN은 None으로 (numeric 열에 'NaN'이 저장되지 않게)."""
    return {c: _norm_val(rec[c]) for c in cols}


def _next_line_numbers(used):
    n = max([int(m.group(1)) for u in used if u and (m := LINE_NO_RE.match(u))] + [0])
    while True:
        n += 1
        yield f"L{n:04d}"


def import_e2e(file, dry_run=False):
    """엑셀을 DB에 반영(dry_run이면 계산만). 반환: 결과 보고 dict."""
    df = read_e2e_excel(file)
    conn = get_db_connection()
    engine = conn.session.get_bind()
    with engine.begin() as c:
        db_lines = pd.read_sql(text("select * from purchase_lines"), c)
        db_rcpt = pd.read_sql(text("select * from purchase_receipts"), c)

    excel_lines = db_lines[db_lines["source"] == "excel"]
    by_line_no = dict(zip(excel_lines["line_no"], excel_lines["id"]))
    by_fp = dict(zip(excel_lines["fingerprint"], excel_lines["id"]))
    dup_in_file = {k for k, v in Counter(df["excel_line_no"].dropna()).items() if v > 1}
    used_line_nos = set(db_lines["line_no"]) | set(df["excel_line_no"].dropna())
    gen = _next_line_numbers(used_line_nos)
    taken = set(db_lines["line_no"])

    matched_ids, inserts, updates, renumbered = set(), [], [], []
    current = db_lines.set_index("id")
    records = [{k: _norm_val(v) for k, v in rec.items()} for rec in df.to_dict("records")]
    for rec in records:
        raw_no = rec["excel_line_no"]
        line_id = None
        if raw_no and raw_no not in dup_in_file and by_line_no.get(raw_no) not in (None, *matched_ids):
            line_id = by_line_no[raw_no]
        elif by_fp.get(rec["fingerprint"]) is not None and by_fp[rec["fingerprint"]] not in matched_ids:
            line_id = by_fp[rec["fingerprint"]]

        if line_id is not None:
            matched_ids.add(line_id)
            old = current.loc[line_id]
            changed = {f: rec[f] for f in LINE_FIELDS if _norm_val(old[f]) != _norm_val(rec[f])}
            if changed:
                updates.append((line_id, changed))
            rec["_line_id"] = line_id
        else:
            line_no = raw_no if raw_no and raw_no not in dup_in_file and raw_no not in taken else next(gen)
            if line_no != raw_no:
                renumbered.append({"엑셀 관리번호": raw_no or "(없음)", "새 관리번호": line_no,
                                   "PO": rec["po_no"], "품명": rec["item_name"]})
            taken.add(line_no)
            rec["line_no"] = line_no
            inserts.append(rec)

    app_rcpt_lines = set(db_rcpt.loc[db_rcpt["source"] == "app", "line_id"])
    excel_rcpt = {r.line_id: r for r in db_rcpt[db_rcpt["source"] == "excel"].itertuples()}
    rcpt_changes, kept_app = [], 0
    for rec in records:
        line_id = rec.get("_line_id")
        want = (_norm_val(rec["received_qty"]) or 0) > 0
        if line_id is not None and line_id in app_rcpt_lines:
            kept_app += 1
            continue
        old = excel_rcpt.get(line_id) if line_id is not None else None
        if old is not None and want and (
            _norm_val(old.qty) == _norm_val(rec["received_qty"])
            and _norm_val(old.received_on) == _norm_val(rec["received_on"])
            and _norm_val(old.location) == _norm_val(rec["destination"])
        ):
            continue
        if old is None and not want:
            continue
        rcpt_changes.append((rec, line_id))

    report = {
        "rows": len(df),
        "new_lines": len(inserts),
        "updated_lines": len(updates),
        "unchanged_lines": len(df) - len(inserts) - len(updates),
        "renumbered": renumbered,
        "receipt_changes": len(rcpt_changes),
        "kept_app_receipts": kept_app,
        "db_only_lines": int((~db_lines["id"].isin(matched_ids)).sum()),
        "changed_fields": Counter(f for _, ch in updates for f in ch),
    }
    if dry_run:
        return report

    with engine.begin() as c:
        new_ids = {}
        for rec in inserts:
            cols = ["line_no"] + LINE_FIELDS
            res = c.execute(
                text(f"insert into purchase_lines ({', '.join(cols)}, source) "
                     f"values ({', '.join(':' + x for x in cols)}, 'excel') returning id"),
                _params(rec, cols),
            )
            new_ids[rec["fingerprint"]] = res.scalar_one()
        for line_id, changed in updates:
            sets = ", ".join(f"{f} = :{f}" for f in changed)
            c.execute(text(f"update purchase_lines set {sets}, updated_at = now() where id = :id"),
                      {**_params(changed, changed), "id": int(line_id)})
        for rec, line_id in rcpt_changes:
            line_id = int(line_id) if line_id is not None else new_ids[rec["fingerprint"]]
            c.execute(text("delete from purchase_receipts where line_id = :l and source = 'excel'"), {"l": line_id})
            p = _params(rec, ["received_qty", "received_on", "destination"])
            if (p["received_qty"] or 0) > 0:
                c.execute(
                    text("insert into purchase_receipts (line_id, qty, received_on, location, source, entered_by) "
                         "values (:l, :q, :d, :loc, 'excel', '엑셀 가져오기')"),
                    {"l": line_id, "q": p["received_qty"], "d": p["received_on"], "loc": p["destination"]},
                )
        sync_vendors_from_lines(c)
    return report


def sync_vendors_from_lines(c):
    """purchase_lines의 공급사 중 거래처 목록에 없는 곳을 추가 (내·외자는 그 공급사 줄의 다수 값)."""
    c.execute(text(
        """
        insert into vendors (name, vendor_type, note)
        select supplier, mode() within group (order by supply_type), '구매 E2E에서 자동 추가'
        from purchase_lines
        where supplier is not null and supply_type is not null
        group by supplier
        on conflict (name) do nothing
        """
    ))
