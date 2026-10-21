"""PO 한 건을 회사 발주서 엑셀 양식(PO 시트)에 채워 xlsx bytes로 만든다.

양식 구조 (0. 나인봇 부품 발주서 양식): B11 발주일, H11 PO번호, 12~15행 병합된 빈칸,
17행 헤더, 18~52행 품목(C 부품번호 · D 한글명 · E 영문명 · F MPQ · G Pack · H 수량 ·
I 단가 · J 금액), 53행 합계.

수식 대신 계산된 값을 채운다 — openpyxl로 쓴 수식은 결과값이 저장되지 않아 휴대폰·메신저
미리보기에서 빈칸으로 보이기 때문. 그리고 PO 시트만 남기고 나머지 시트(단가표 RAWDATA 등)는
지운다 — 다른 거래처에 보내는 발주서에 단가표가 딸려가면 안 되므로."""

import io
from copy import copy
from datetime import date, datetime

import openpyxl

from po_catalog import PO_SHEET

FIRST_ROW = 18
TEMPLATE_LAST_ROW = 52
TOTAL_ROW = 53
COLS = "BCDEFGHIJ"
CURRENCY_FORMATS = {
    "KRW": '#,##0"원"',
    "USD": None,  # 양식 기본($) 서식 유지
}


def _as_datetime(d):
    if d is None:
        return None
    if isinstance(d, datetime):
        return d
    if isinstance(d, date):
        return datetime(d.year, d.month, d.day)
    return d


def build_po_xlsx(header, lines, template_bytes):
    """header: {po_number, po_date, due_date, currency}
    lines: [{package_no, name_ko, desc_en, mpq, qty, unit_price}, ...]
    template_bytes: load_latest_template()[2] — st.download_button의 지연 생성 콜백이 별도
    스레드에서 돌기 때문에, DB/캐시 조회는 호출하는 쪽에서 미리 해서 넘겨받는다."""
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb[PO_SHEET]
    currency = header.get("currency") or "USD"

    ws["B11"] = _as_datetime(header["po_date"])
    ws["H11"] = header["po_number"]
    if header.get("due_date"):
        ws["B12"] = "Required Delivery Date"
        ws["F12"] = _as_datetime(header["due_date"])
        ws["F12"].number_format = "yyyy-mm-dd"

    # 35줄을 넘으면 합계 행 앞에 줄을 끼워 넣고 마지막 품목 줄 서식을 복사한다.
    extra = max(0, len(lines) - (TEMPLATE_LAST_ROW - FIRST_ROW + 1))
    if extra:
        ws.insert_rows(TOTAL_ROW, extra)
        for r in range(TOTAL_ROW, TOTAL_ROW + extra):
            ws.row_dimensions[r].height = ws.row_dimensions[TEMPLATE_LAST_ROW].height
            for col in COLS:
                ws[f"{col}{r}"]._style = copy(ws[f"{col}{TEMPLATE_LAST_ROW}"]._style)
    last_row = TEMPLATE_LAST_ROW + extra
    total_row = TOTAL_ROW + extra

    for r in range(FIRST_ROW, last_row + 1):
        for col in "CDEFGHIJ":
            ws[f"{col}{r}"] = None

    total_qty = 0
    total_amt = 0.0
    for i, ln in enumerate(lines):
        r = FIRST_ROW + i
        qty = float(ln["qty"])
        price = float(ln["unit_price"]) if ln.get("unit_price") not in (None, "") else None
        mpq = float(ln["mpq"]) if ln.get("mpq") else None
        amt = round(qty * price, 4) if price is not None else None
        ws[f"C{r}"] = ln.get("package_no")
        ws[f"D{r}"] = ln.get("name_ko")
        ws[f"E{r}"] = ln.get("desc_en")
        ws[f"F{r}"] = mpq
        ws[f"G{r}"] = round(qty / mpq, 2) if mpq else None
        ws[f"H{r}"] = qty
        ws[f"I{r}"] = price
        ws[f"J{r}"] = amt
        total_qty += qty
        total_amt += amt or 0

    ws[f"G{total_row}"] = "합계"
    ws[f"H{total_row}"] = total_qty
    ws[f"J{total_row}"] = total_amt
    ws["J17"] = f"Amount({currency})"

    fmt = CURRENCY_FORMATS.get(currency, "#,##0.00")
    if fmt:
        for r in list(range(FIRST_ROW, last_row + 1)) + [total_row]:
            ws[f"I{r}"].number_format = fmt
            ws[f"J{r}"].number_format = fmt

    for name in list(wb.sheetnames):
        if name != PO_SHEET:
            del wb[name]
    # 지운 시트를 가리키는 이름 정의가 남아 있으면 엑셀이 "복구" 경고를 띄우므로 같이 정리한다.
    for dn in list(wb.defined_names):
        if "!" in str(wb.defined_names[dn].value) and PO_SHEET not in str(wb.defined_names[dn].value):
            del wb.defined_names[dn]
    wb.active = 0

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
