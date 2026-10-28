"""발주서(PO) 엑셀 양식과 그 안의 부품 카탈로그 관리.
양식에 거래처 단가가 들어있어서 공개 저장소 대신 DB(po_templates/parts_catalog)에 보관하고,
관리자가 앱에서 양식 엑셀을 올리면 RAWDATA 시트를 파싱해 카탈로그를 통째로 교체한다."""

import io
import re

import openpyxl
import pandas as pd
import streamlit as st
from sqlalchemy import text

from db import get_db_connection

PO_SHEET = "PO"
RAWDATA_PREFIX = "부품 정보 RAWDATA"
BOXHERO_SHEET = "박스히어로"


def _norm(h):
    return re.sub(r"\s+", " ", str(h or "")).strip()


def _find_rawdata_sheet(wb):
    for name in wb.sheetnames:
        if name.startswith(RAWDATA_PREFIX):
            return wb[name]
    return None


def _price_col(headers, tier):
    """'(1-499)' / '(500-5000)' 단가 열. 연도가 붙은 최신 단가("26년 Price")를 우선한다."""
    cands = [i for i, h in enumerate(headers) if "price" in h.lower() and tier in h]
    with_year = [i for i in cands if "년" in headers[i]]
    return (with_year or cands or [None])[-1]


def parse_catalog(wb):
    ws = _find_rawdata_sheet(wb)
    if ws is None:
        raise ValueError(f"'{RAWDATA_PREFIX}…' 시트를 찾을 수 없어요.")
    rows = list(ws.iter_rows(values_only=True))
    headers = [_norm(h) for h in rows[0]]

    def col(name):
        return headers.index(name) if name in headers else None

    idx = {
        "package_no": col("Package No."),
        "sku": col("SKU"),
        "desc_en": col("English description"),
        "name_ko": col("Korean description"),
        "mpq": col("MPQ"),
        "category": col("Category"),
        "price_low": _price_col(headers, "(1-499)"),
        "price_high": _price_col(headers, "(500-5000)"),
    }
    if idx["package_no"] is None:
        raise ValueError("RAWDATA 시트에 'Package No.' 열이 없어요.")

    # 양식의 PO 시트가 한글명을 박스히어로 시트(바코드 -> 제품명)에서 가져오므로 그걸 우선한다.
    bh_names = {}
    if BOXHERO_SHEET in wb.sheetnames:
        for code, name in wb[BOXHERO_SHEET].iter_rows(min_row=2, max_col=2, values_only=True):
            if code and name:
                bh_names[str(code).strip()] = str(name).strip()

    def val(row, key):
        i = idx[key]
        return row[i] if i is not None and i < len(row) else None

    def num(v):
        try:
            return float(v) if v not in (None, "") else None
        except (TypeError, ValueError):
            return None

    records = {}
    for row in rows[1:]:
        pkg = val(row, "package_no")
        if not pkg or not str(pkg).strip():
            continue
        pkg = str(pkg).strip()
        if pkg in records:
            continue
        sku = val(row, "sku")
        records[pkg] = {
            "package_no": pkg,
            "sku": str(sku).strip() if sku else None,
            "name_ko": bh_names.get(pkg) or (str(val(row, "name_ko")).strip() if val(row, "name_ko") else None),
            "desc_en": str(val(row, "desc_en")).strip() if val(row, "desc_en") else None,
            "mpq": num(val(row, "mpq")),
            "price_low": num(val(row, "price_low")),
            "price_high": num(val(row, "price_high")),
            "category": str(val(row, "category")).strip() if val(row, "category") else None,
        }
    if not records:
        raise ValueError("RAWDATA 시트에서 부품을 하나도 읽지 못했어요.")
    return pd.DataFrame(records.values())


def save_po_template(file_bytes, file_name, uploaded_by):
    """양식을 검증하고 저장한 뒤 카탈로그를 통째로 교체. 반환: 카탈로그 품목 수."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    if PO_SHEET not in wb.sheetnames:
        raise ValueError(f"'{PO_SHEET}' 시트가 없어요. 발주서 양식 파일이 맞는지 확인해주세요.")
    catalog = parse_catalog(wb)

    conn = get_db_connection()
    engine = conn.session.get_bind()
    with engine.begin() as c:
        c.execute(
            text("insert into po_templates (file_name, file_bytes, uploaded_by) values (:n, :b, :u)"),
            {"n": file_name, "b": file_bytes, "u": uploaded_by},
        )
        c.execute(text("delete from parts_catalog"))
        catalog.to_sql("parts_catalog", con=c, if_exists="append", index=False, method="multi", chunksize=500)
    load_latest_template.clear()
    load_catalog.clear()
    template_last_po_number.clear()
    return len(catalog)


@st.cache_data(ttl=600, show_spinner=False)
def load_latest_template():
    """(file_name, uploaded_at, file_bytes) 또는 None."""
    conn = get_db_connection()
    engine = conn.session.get_bind()
    with engine.connect() as c:
        row = c.execute(
            text("select file_name, uploaded_at, file_bytes from po_templates order by uploaded_at desc limit 1")
        ).first()
    if row is None:
        return None
    return row[0], row[1], bytes(row[2])


@st.cache_data(ttl=600, show_spinner=False)
def load_catalog():
    conn = get_db_connection()
    return conn.query("select * from parts_catalog order by package_no", ttl=0)


def search_catalog(q, limit=20):
    df = load_catalog()
    if df.empty or not q.strip():
        return df.head(0)
    q = q.strip().lower()
    hay = (
        df["package_no"].fillna("") + " " + df["sku"].fillna("") + " "
        + df["name_ko"].fillna("") + " " + df["desc_en"].fillna("")
    ).str.lower()
    return df[hay.str.contains(q, regex=False)].head(limit)


@st.cache_data(ttl=600, show_spinner=False)
def template_last_po_number():
    """양식 H11에 마지막으로 적혀 있던 PO번호 — 다음 PO번호 계산에 참고 (procurement.next_po_number)."""
    tpl = load_latest_template()
    if tpl is None:
        return None
    return openpyxl.load_workbook(io.BytesIO(tpl[2]))[PO_SHEET]["H11"].value
